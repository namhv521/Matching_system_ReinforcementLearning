"""Excel Builder for KLTN RL Matching Checklist and Roadmap."""
import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from scripts.kltn_excel_data import WEEKS_DATA, STEPS_7_DATA, QA_WEEK1_DATA, ALGO_COMP_DATA, TASKS_56

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "KE_HOACH_CHECKLIST_KLTN_RL_MATCHING.xlsx"

# Palettes
NAVY_HEADER = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
STEEL_HEADER = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
LIGHT_BLUE = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
CARD_BG = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")

FILL_DONE = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FONT_DONE = Font(name="Calibri", size=10, color="006100", bold=True)
FILL_DOING = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
FONT_DOING = Font(name="Calibri", size=10, color="9C6500", bold=True)
FILL_TODO = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")
FONT_TODO = Font(name="Calibri", size=10, color="595959", bold=False)
FILL_REVIEW = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FONT_REVIEW = Font(name="Calibri", size=10, color="9C0006", bold=True)

FONT_TITLE = Font(name="Calibri", size=15, bold=True, color="1F4E78")
FONT_SUBTITLE = Font(name="Calibri", size=11, italic=True, color="595959")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="000000")
FONT_REGULAR = Font(name="Calibri", size=10, color="000000")

THIN_GRAY = Side(style="thin", color="D9D9D9")
MEDIUM_BLUE = Side(style="medium", color="1F4E78")
BORDER_CELL = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
BORDER_HEADER = Border(left=THIN_GRAY, right=THIN_GRAY, top=MEDIUM_BLUE, bottom=MEDIUM_BLUE)

def style_header_row(ws, row_idx, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = NAVY_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_HEADER

def style_status_cell(cell, status_val):
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if status_val == "Hoàn thành":
        cell.fill = FILL_DONE; cell.font = FONT_DONE
    elif status_val == "Đang thực hiện":
        cell.fill = FILL_DOING; cell.font = FONT_DOING
    elif status_val == "Chưa bắt đầu":
        cell.fill = FILL_TODO; cell.font = FONT_TODO
    else:
        cell.fill = FILL_REVIEW; cell.font = FONT_REVIEW

def autofit(ws, max_cap=65):
    ws.views.sheetView[0].showGridLines = True
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            lines = str(cell.value or "").split("\n")
            line_len = max(len(l) for l in lines) if lines else 0
            if cell.font and cell.font.size and cell.font.size > 12:
                line_len = int(line_len * 1.3)
            if line_len > max_len:
                max_len = line_len
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), max_cap)

def build_sheet1_dashboard(wb):
    ws = wb.create_sheet(title="Dashboard_TongQuan")
    ws.cell(row=2, column=2, value="KẾ HOẠCH TỔNG THỂ & DASHBOARD THEO DÕI KHÓA LUẬN TỐT NGHIỆP").font = FONT_TITLE
    ws.cell(row=3, column=2, value="Đề tài: Hệ thống Reinforcement Learning Tối ưu Phân bổ Sinh viên – Giảng viên Hướng dẫn | Sinh viên: Hoàng Văn Nam (11236160)").font = FONT_SUBTITLE
    ws.cell(row=4, column=2, value="Thời gian thực hiện: 21/09/2026 – 15/11/2026 (8 tuần / 56 ngày) | Deadline bảo vệ/nộp: 15/11/2026").font = Font(name="Calibri", size=10, bold=True, color="C00000")

    kpis = [
        ("Tổng thời gian", "8 Tuần (56 Ngày)", "21/09 -> 15/11/2026"),
        ("Mô hình chính", "Maskable PPO", "PPO có Action Masking"),
        ("Mô hình đối chứng", "DQN / QR-DQN / A2C", "Đối sánh Unmasked RL"),
        ("Baseline bắt buộc", "Hungarian, SPA, Greedy", "Hungarian là Upper Bound"),
        ("Mục tiêu Quota", "0 vi phạm (100% hợp lệ)", "Ràng buộc cứng"),
        ("Mục tiêu Compatibility", ">= 92% của Hungarian", "Validation & Test split"),
    ]
    for idx, (title, val, note) in enumerate(kpis):
        col = 2 + idx * 2
        ws.merge_cells(start_row=6, start_column=col, end_row=6, end_column=col+1)
        ws.merge_cells(start_row=7, start_column=col, end_row=7, end_column=col+1)
        ws.merge_cells(start_row=8, start_column=col, end_row=8, end_column=col+1)
        c_title = ws.cell(row=6, column=col, value=title)
        c_title.font = Font(name="Calibri", size=9, bold=True, color="595959")
        c_title.fill = LIGHT_BLUE; c_title.alignment = Alignment(horizontal="center", vertical="center")
        c_val = ws.cell(row=7, column=col, value=val)
        c_val.font = Font(name="Calibri", size=12, bold=True, color="1F4E78")
        c_val.fill = CARD_BG; c_val.alignment = Alignment(horizontal="center", vertical="center")
        c_note = ws.cell(row=8, column=col, value=note)
        c_note.font = Font(name="Calibri", size=8, italic=True, color="7F7F7F")
        c_note.fill = CARD_BG; c_note.alignment = Alignment(horizontal="center", vertical="center")

    ws.cell(row=10, column=2, value="LỘ TRÌNH 8 TUẦN THEO 7 BƯỚC THIẾT KẾ HỆ THỐNG MACHINE LEARNING").font = Font(name="Calibri", size=13, bold=True, color="1F4E78")
    headers = ["Tuần", "Khoảng thời gian", "Giai đoạn ML (7 Bước)", "Trọng tâm kỹ thuật & Deliverables", "Câu hỏi then chốt giải quyết", "Tiêu chí nghiệm thu (Gate)", "Trạng thái"]
    for c_idx, h in enumerate(headers, start=2):
        ws.cell(row=12, column=c_idx, value=h)
    style_header_row(ws, 12, len(headers)+1)

    for idx, row_data in enumerate(WEEKS_DATA, start=13):
        ws.row_dimensions[idx].height = 28
        for c_idx, val in enumerate(row_data, start=2):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = FONT_REGULAR; cell.border = BORDER_CELL
            if c_idx == 2:
                cell.font = FONT_BOLD; cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 3:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 8:
                style_status_cell(cell, val)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if idx % 2 == 1: cell.fill = ZEBRA_FILL
    autofit(ws)

def build_sheet2_checklist(wb):
    ws = wb.create_sheet(title="Checklist_HangNgay")
    ws.cell(row=2, column=1, value="BẢNG CHECKLIST CÔNG VIỆC HÀNG NGÀY (21/09/2026 – 15/11/2026)").font = FONT_TITLE
    ws.cell(row=3, column=1, value="Kế hoạch tác chiến 56 ngày chia theo 7 bước thiết kế hệ thống Machine Learning & Reinforcement Learning").font = FONT_SUBTITLE

    headers = [
        "STT", "Ngày", "Thứ", "Tuần", "Giai đoạn ML",
        "Hạng mục công việc chi tiết (Daily Task)",
        "Mục tiêu cụ thể cần đạt",
        "Deliverables / File sản phẩm",
        "Lệnh kiểm tra / Test command",
        "Người phụ trách", "Trạng thái", "Ghi chú & Rủi ro"
    ]
    for c_idx, h in enumerate(headers, start=1):
        ws.cell(row=5, column=c_idx, value=h)
    style_header_row(ws, 5, len(headers))

    start_date = datetime.date(2026, 9, 21)
    days_vn = ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật"]
    cur_date = start_date

    for idx, item in enumerate(TASKS_56, start=1):
        row_num = idx + 5
        ws.row_dimensions[row_num].height = 24
        stage, task, obj, deliverable, cmd, owner, status, note = item
        day_vn = days_vn[cur_date.weekday()]
        week_num = f"Tuần {(idx - 1) // 7 + 1}"
        date_str = cur_date.strftime("%d/%m/%Y")
        row_vals = [idx, date_str, day_vn, week_num, stage, task, obj, deliverable, cmd, owner, status, note]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_num, column=c_idx, value=val)
            cell.font = FONT_REGULAR; cell.border = BORDER_CELL
            if c_idx in (1, 2, 3, 4, 10):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if c_idx == 1: cell.font = FONT_BOLD
            elif c_idx == 11:
                style_status_cell(cell, val)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if idx % 2 == 1: cell.fill = ZEBRA_FILL
        cur_date += datetime.timedelta(days=1)
    autofit(ws)
def build_sheet3_weekly(wb):
    ws = wb.create_sheet(title="KeHoach_HangTuan")
    ws.cell(row=2, column=1, value="CHI TIẾT KẾ HOẠCH HÀNG TUẦN (8 TUẦN TÁC CHIẾN)").font = FONT_TITLE
    ws.cell(row=3, column=1, value="Mục tiêu chiến lược, câu hỏi nghiên cứu giải quyết, deliverables và tiêu chí nghiệm thu từng tuần").font = FONT_SUBTITLE

    headers = [
        "Tuần", "Khoảng thời gian", "Giai đoạn ML (7 Bước)",
        "Trọng tâm kỹ thuật & Deliverables",
        "Các câu hỏi nghiên cứu cốt lõi cần giải quyết",
        "Tiêu chí nghiệm thu (Promotion / Acceptance Gate)", "Trạng thái"
    ]
    for c_idx, h in enumerate(headers, start=1):
        ws.cell(row=5, column=c_idx, value=h)
    style_header_row(ws, 5, len(headers))

    for idx, row_data in enumerate(WEEKS_DATA, start=6):
        ws.row_dimensions[idx].height = 45
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = FONT_REGULAR; cell.border = BORDER_CELL
            if c_idx == 1:
                cell.font = FONT_BOLD; cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 2:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 7:
                style_status_cell(cell, val)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if idx % 2 == 1: cell.fill = ZEBRA_FILL
    autofit(ws)

def build_sheet4_steps7(wb):
    ws = wb.create_sheet(title="ThietKe_RL_7Buoc")
    ws.cell(row=2, column=1, value="ÁNH XẠ 7 BƯỚC THIẾT KẾ HỆ THỐNG MACHINE LEARNING CHO BÀI TOÁN PHÂN BỔ RL").font = FONT_TITLE
    ws.cell(row=3, column=1, value="Tuân thủ chuẩn Machine Learning System Design & MLE Workflow cho đề tài Khóa luận tốt nghiệp").font = FONT_SUBTITLE

    headers = [
        "Bước", "Tên bước thiết kế hệ thống ML",
        "Bản chất bài toán trong đề tài KLTN",
        "Đầu vào (Inputs)", "Thành phần kỹ thuật & Thuật toán áp dụng",
        "Đầu ra kỹ thuật (Artifacts / Outputs)",
        "Tiêu chí chất lượng & Quality Gates", "Tuần thực hiện"
    ]
    for c_idx, h in enumerate(headers, start=1):
        ws.cell(row=5, column=c_idx, value=h)
    style_header_row(ws, 5, len(headers))

    for idx, row_data in enumerate(STEPS_7_DATA, start=6):
        ws.row_dimensions[idx].height = 55
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = FONT_REGULAR; cell.border = BORDER_CELL
            if c_idx == 1:
                cell.font = FONT_BOLD; cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in (2, 8):
                cell.font = FONT_BOLD if c_idx == 2 else FONT_REGULAR
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if idx % 2 == 1: cell.fill = ZEBRA_FILL
    autofit(ws)

def build_sheet5_qa_week1(wb):
    ws = wb.create_sheet(title="KienTruc_Feature_Reward_Loss")
    ws.cell(row=2, column=1, value="GIẢI TRÌNH CHI TIẾT 6 CÂU HỎI CỐT LÕI (WEEK 1 RESEARCH SPECIFICATION)").font = FONT_TITLE
    ws.cell(row=3, column=1, value="Đặc tả kỹ thuật sâu về Features, Loss, Reward, Đánh giá, Vòng đời Multi-turn MDP vs Hồi quy, và Đối chuẩn").font = FONT_SUBTITLE

    headers = ["STT", "Câu hỏi nghiên cứu cốt lõi", "Thành phần kỹ thuật", "Chi tiết đặc tả & Công thức toán học", "Vai trò trong hệ thống", "Phương thức kiểm chứng thực nghiệm"]
    for c_idx, h in enumerate(headers, start=1):
        ws.cell(row=5, column=c_idx, value=h)
    style_header_row(ws, 5, len(headers))

    for idx, row_data in enumerate(QA_WEEK1_DATA, start=6):
        ws.row_dimensions[idx].height = 65
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = FONT_REGULAR; cell.border = BORDER_CELL
            if c_idx == 1:
                cell.font = FONT_BOLD; cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in (2, 3):
                cell.font = FONT_BOLD; cell.alignment = Alignment(vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if idx % 2 == 1: cell.fill = ZEBRA_FILL
    autofit(ws)

def build_sheet6_comparison(wb):
    ws = wb.create_sheet(title="SoSanh_MoHinh_RL_Baselines")
    ws.cell(row=2, column=1, value="BẢNG ĐỐI SÁNH NĂNG LỰC CÁC THUẬT TOÁN (RL MODELS VS BASELINES)").font = FONT_TITLE
    ws.cell(row=3, column=1, value="So sánh lý thuyết và thực nghiệm giữa các họ thuật toán Reinforcement Learning và thuật toán cổ điển").font = FONT_SUBTITLE

    headers = [
        "Họ giải thuật", "Thuật toán", "Cơ chế hoạt động chính",
        "Hỗ trợ Ràng buộc cứng (Hard Quota)", "Tính tối ưu trong Batch",
        "Tính thích ứng Online Streaming", "Độ phức tạp thời gian",
        "Kết quả thực nghiệm trên Validation", "Vai trò trong Khóa luận"
    ]
    for c_idx, h in enumerate(headers, start=1):
        ws.cell(row=5, column=c_idx, value=h)
    style_header_row(ws, 5, len(headers))

    for idx, row_data in enumerate(ALGO_COMP_DATA, start=6):
        ws.row_dimensions[idx].height = 45
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=idx, column=c_idx, value=val)
            cell.font = FONT_REGULAR; cell.border = BORDER_CELL
            if c_idx in (1, 2):
                cell.font = FONT_BOLD; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif c_idx in (4, 5, 6, 7):
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if idx % 2 == 1: cell.fill = ZEBRA_FILL
    autofit(ws)

def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    print("Building Sheet 1: Dashboard_TongQuan...")
    build_sheet1_dashboard(wb)
    print("Building Sheet 2: Checklist_HangNgay...")
    build_sheet2_checklist(wb)
    print("Building Sheet 3: KeHoach_HangTuan...")
    build_sheet3_weekly(wb)
    print("Building Sheet 4: ThietKe_RL_7Buoc...")
    build_sheet4_steps7(wb)
    print("Building Sheet 5: KienTruc_Feature_Reward_Loss...")
    build_sheet5_qa_week1(wb)
    print("Building Sheet 6: SoSanh_MoHinh_RL_Baselines...")
    build_sheet6_comparison(wb)

    wb.save(EXCEL_PATH)
    print(f"SUCCESS: Workbook successfully generated at: {EXCEL_PATH}")

if __name__ == "__main__":
    main()
